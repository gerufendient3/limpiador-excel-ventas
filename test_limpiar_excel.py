"""Pruebas independientes para limpiar_excel.py; no importa el generador."""

import hashlib
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SCRIPT = Path(__file__).with_name("limpiar_excel.py")


def sha256(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def test_limpieza_independiente(tmp_path):
    entrada = tmp_path / "entrada_manual.xlsx"
    filas = [
        {"nombre": "josé", "fecha": "13/10/2026", "monto": 125, "ciudad": " Lima "},
        {"nombre": "josé", "fecha": "13/10/2026", "monto": 125, "ciudad": " Lima "},  # duplicado exacto
        {"nombre": " josé", "fecha": "13/10/2026", "monto": 125, "ciudad": " Lima "},  # distinto al comparar original
        {"nombre": "Ana", "fecha": "2026-02-30", "monto": 20, "ciudad": "Ica"},
        {"nombre": "Luis", "fecha": "2026-10-13", "monto": "mil", "ciudad": "Cusco"},
        {"nombre": "", "fecha": "13-10-26", "monto": -4, "ciudad": "Piura"},
    ]
    pd.DataFrame(filas, columns=["nombre", "fecha", "monto", "ciudad"]).to_excel(
        entrada, sheet_name="ventas", index=False
    )
    hash_antes = sha256(entrada)

    resultado = subprocess.run(
        [sys.executable, str(SCRIPT), str(entrada)], capture_output=True, text=True, check=False
    )
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    assert "Duplicados exactos eliminados: 1" in resultado.stdout
    assert "Fechas inválidas: 1" in resultado.stdout
    assert "Nombres vacíos: 1" in resultado.stdout
    assert "Montos con error: 2" in resultado.stdout
    assert sha256(entrada) == hash_antes, "El archivo original fue modificado."

    limpia = tmp_path / "ventas_limpias.xlsx"
    informe = tmp_path / "informe_errores.xlsx"
    assert limpia.exists() and informe.exists()
    df = pd.read_excel(limpia, sheet_name="ventas", dtype=object, keep_default_na=False)
    assert len(df) == 5
    assert "José" in df["nombre"].tolist()
    assert " josé" not in df["nombre"].tolist()
    assert df.loc[df["nombre"] == "Ana", "fecha"].iloc[0] == ""
    assert df.loc[df["nombre"] == "Luis", "monto"].iloc[0] == ""
    # La fila con espacio inicial no se consideró duplicada antes de normalizar.
    assert len(df[df["nombre"] == "José"]) == 2

    libro = load_workbook(informe, data_only=True)
    assert libro.sheetnames == ["incidencias", "resumen"]
    incidencias = list(libro["incidencias"].values)
    cabecera = incidencias[0]
    registros = [dict(zip(cabecera, fila)) for fila in incidencias[1:]]
    tipos = [r["tipo_error"] for r in registros]
    assert "duplicado exacto eliminado" in tipos
    assert "fecha inválida" in tipos
    assert "monto no numérico" in tipos
    assert "monto negativo" in tipos
    assert "nombre vacío" in tipos
    libro.close()


def test_faltan_columnas_da_error_claro(tmp_path):
    entrada = tmp_path / "incompleto.xlsx"
    pd.DataFrame([{"nombre": "Ana"}]).to_excel(entrada, sheet_name="ventas", index=False)
    resultado = subprocess.run(
        [sys.executable, str(SCRIPT), str(entrada)], capture_output=True, text=True, check=False
    )
    assert resultado.returncode != 0
    assert "Faltan columnas obligatorias" in resultado.stderr


def test_formatos_fecha_montos_y_particulas(tmp_path):
    entrada = tmp_path / "formatos.xlsx"
    filas = [
        {"nombre": "juan de la cruz", "fecha": "13-10-26", "monto": "S/ 1,250", "ciudad": "Lima"},
        {"nombre": "MARÍA DEL CARMEN", "fecha": "2026-10-13", "monto": 80.5, "ciudad": "Ica"},
    ]
    pd.DataFrame(filas, columns=["nombre", "fecha", "monto", "ciudad"]).to_excel(
        entrada, sheet_name="ventas", index=False
    )
    resultado = subprocess.run(
        [sys.executable, str(SCRIPT), str(entrada)], capture_output=True, text=True, check=False
    )
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    assert "Montos con error: 0" in resultado.stdout

    df = pd.read_excel(tmp_path / "ventas_limpias.xlsx", sheet_name="ventas", dtype=object, keep_default_na=False)
    assert df.loc[0, "nombre"] == "Juan de la Cruz"
    assert df.loc[1, "nombre"] == "María del Carmen"
    assert df.loc[0, "fecha"] == "2026-10-13"
    assert float(df.loc[0, "monto"]) == 1250
    assert float(df.loc[1, "monto"]) == 80.5
