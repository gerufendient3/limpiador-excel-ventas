"""Genera un Excel reproducible con datos de ventas de prueba, incluidos errores intencionales."""

from pathlib import Path
import random
import re
from datetime import date, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SEMILLA = 20261013
ARCHIVO_SALIDA = Path(__file__).resolve().parent / "datos_prueba.xlsx"
COLUMNAS = ["nombre", "fecha", "monto", "ciudad"]

NOMBRES = [
    "María Fernanda Torres", "José Luis Quispe", "Ana Lucía Mendoza",
    "Diego Ramírez Huamán", "Valeria Flores", "Carlos Alberto Rojas",
    "Lucía Vargas", "Andrés Castillo", "Camila Chávez", "Miguel Ángel Paredes",
    "Daniela Salazar", "Jorge Luis Gutiérrez", "Sofía Herrera", "Renato Cárdenas",
    "Paola Fernández", "Luis Enrique Medina", "Gabriela Ramos", "Óscar Villanueva",
    "Rosa Milagros Soto", "Fernando Aguilar", "Patricia León", "Marco Antonio Silva",
    "Andrea Espinoza", "Pedro Pablo Navarro", "Claudia Reyes", "Bruno Cabrera",
    "Milagros Poma", "Alonso Valdez", "Fiorella Acuña", "Hugo Carrillo",
]
CIUDADES = ["Lima", "Arequipa", "Cusco", "Trujillo", "Chiclayo", "Piura", "Ica", "Huancayo"]


def formato_fecha(fecha: date, formato: int) -> str:
    """Devuelve una fecha textual en uno de los tres formatos requeridos."""
    if formato == 0:
        return fecha.strftime("%d/%m/%Y")
    if formato == 1:
        return fecha.strftime("%Y-%m-%d")
    return fecha.strftime("%d-%m-%y")


def generar_datos() -> pd.DataFrame:
    rng = random.Random(SEMILLA)
    filas = []
    inicio = date(2025, 1, 1)

    # Crear 45 filas base con fechas distribuidas entre los tres formatos.
    for i in range(45):
        fecha = inicio + timedelta(days=rng.randint(0, 620))
        filas.append({
            "nombre": rng.choice(NOMBRES),
            "fecha": formato_fecha(fecha, i % 3),
            "monto": round(rng.uniform(10, 5000), 2),
            "ciudad": rng.choice(CIUDADES),
        })

    # Inyectar exactamente tres fechas inválidas en filas distintas.
    filas[0]["fecha"] = "31/02/2026"
    filas[1]["fecha"] = "texto"
    filas[2]["fecha"] = "45/13/2026"

    # Inyectar exactamente cuatro nombres vacíos.
    for indice in (3, 4, 5, 6):
        filas[indice]["nombre"] = ""

    # Casos de espacios sobrantes y de uso de mayúsculas/minúsculas.
    filas[7]["nombre"] = "  María Fernanda Torres"
    filas[8]["nombre"] = "José Luis Quispe  "
    filas[9]["nombre"] = "ana lucía mendoza"
    filas[10]["nombre"] = "CARLOS ALBERTO ROJAS"

    df_base = pd.DataFrame(filas, columns=COLUMNAS)

    # Añadir cinco copias exactas de filas limpias anteriores. Se eligen filas
    # sin fecha inválida ni nombre vacío para no alterar esos conteos exactos.
    indices_copia = [10, 11, 12, 13, 14]
    copias = df_base.iloc[indices_copia].copy(deep=True)
    df = pd.concat([df_base, copias], ignore_index=True)
    return df


def fecha_invalida(valor: str) -> bool:
    """Reconoce los tres valores inválidos insertados deliberadamente."""
    return valor in {"31/02/2026", "texto", "45/13/2026"}


def validar_dataframe(df: pd.DataFrame) -> dict:
    assert len(df) == 50, f"Se esperaban 50 filas; se encontraron {len(df)}."
    assert list(df.columns) == COLUMNAS, f"Columnas incorrectas: {list(df.columns)}"
    duplicadas = int(df.duplicated(keep="first").sum())
    assert duplicadas == 5, f"Se esperaban 5 duplicados adicionales; se encontraron {duplicadas}."

    fechas = df["fecha"].astype(str).tolist()
    patron_dmy = re.compile(r"^\d{2}/\d{2}/\d{4}$")
    patron_ymd = re.compile(r"^\d{4}-\d{2}-\d{2}$")
    patron_dmy2 = re.compile(r"^\d{2}-\d{2}-\d{2}$")
    formatos = {
        "dd/mm/aaaa": sum(bool(patron_dmy.fullmatch(x)) for x in fechas),
        "aaaa-mm-dd": sum(bool(patron_ymd.fullmatch(x)) for x in fechas),
        "dd-mm-aa": sum(bool(patron_dmy2.fullmatch(x)) for x in fechas),
    }
    assert all(n > 0 for n in formatos.values()), "Falta al menos un formato de fecha requerido."

    invalidas = sum(fecha_invalida(x) for x in fechas)
    vacios = int(df["nombre"].map(lambda x: x == "" or pd.isna(x)).sum())
    assert invalidas == 3, f"Se esperaban 3 fechas inválidas; se encontraron {invalidas}."
    assert vacios == 4, f"Se esperaban 4 nombres vacíos; se encontraron {vacios}."

    nombres = df["nombre"].fillna("").astype(str).tolist()
    espacios = sum(n != n.strip() for n in nombres if n.strip())
    minusculas = sum(n == n.lower() and any(c.isalpha() for c in n) for n in nombres if n.strip())
    mayusculas = sum(n == n.upper() and any(c.isalpha() for c in n) for n in nombres if n.strip())
    assert espacios >= 2, "No se encontraron suficientes nombres con espacios sobrantes."
    assert minusculas >= 1, "No se encontró ningún nombre completamente en minúsculas."
    assert mayusculas >= 1, "No se encontró ningún nombre completamente en mayúsculas."
    assert df["monto"].between(10, 5000).all(), "Hay montos fuera del intervalo de 10 a 5000 soles."

    return {
        "filas": len(df), "duplicados": duplicadas, "formatos": formatos,
        "fechas_invalidas": invalidas, "nombres_vacios": vacios,
        "nombres_con_espacios": espacios, "nombres_minusculas": minusculas,
        "nombres_mayusculas": mayusculas,
    }


def guardar_excel(df: pd.DataFrame) -> None:
    # Escribir fechas como texto y montos como números. La hoja se limita a ventas.
    with pd.ExcelWriter(ARCHIVO_SALIDA, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="ventas")
        hoja = writer.sheets["ventas"]
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="24527A")
            celda.alignment = Alignment(horizontal="center")
        hoja.column_dimensions["A"].width = 28
        hoja.column_dimensions["B"].width = 16
        hoja.column_dimensions["C"].width = 14
        hoja.column_dimensions["D"].width = 16
        for fila in range(2, hoja.max_row + 1):
            hoja.cell(fila, 2).number_format = "@"  # Columna fecha como texto.
            hoja.cell(fila, 3).number_format = '"S/ "#,##0.00'


def validar_archivo_guardado(resumen_esperado: dict) -> None:
    """Reabrir con openpyxl y verificar el archivo real, no solo el DataFrame."""
    libro = load_workbook(ARCHIVO_SALIDA, data_only=False)
    assert libro.sheetnames == ["ventas"], f"Hojas inesperadas: {libro.sheetnames}"
    hoja = libro["ventas"]
    assert hoja.max_row == 51, f"Se esperaba cabecera + 50 filas; hay {hoja.max_row} filas de hoja."
    assert [hoja.cell(1, col).value for col in range(1, 5)] == COLUMNAS, "Cabecera u orden de columnas incorrectos."

    # Volver a construir el DataFrame a partir de los valores realmente guardados.
    valores = list(hoja.iter_rows(min_row=2, max_col=4, values_only=True))
    df_guardado = pd.DataFrame(valores, columns=COLUMNAS)
    # Excel puede representar una cadena vacía como una celda vacía (None).
    df_guardado["nombre"] = df_guardado["nombre"].map(lambda x: "" if x is None else x)
    df_guardado["fecha"] = df_guardado["fecha"].map(lambda x: "" if x is None else str(x))
    resumen_guardado = validar_dataframe(df_guardado)
    assert resumen_guardado == resumen_esperado, "Los datos guardados no coinciden con las validaciones previas."

    # Comprobar que las fechas se guardaron como texto y los montos como números.
    for fila in range(2, hoja.max_row + 1):
        fecha_celda = hoja.cell(fila, 2)
        monto_celda = hoja.cell(fila, 3)
        assert isinstance(fecha_celda.value, str), f"La fecha de la fila {fila} no está guardada como texto."
        assert isinstance(monto_celda.value, (int, float)), f"El monto de la fila {fila} no es numérico."
    libro.close()


def main() -> None:
    df = generar_datos()
    resumen = validar_dataframe(df)
    guardar_excel(df)
    validar_archivo_guardado(resumen)

    print("VALIDACIÓN COMPLETADA: datos_prueba.xlsx cumple todos los requisitos.")
    print(f"Filas de datos: {resumen['filas']}")
    print(f"Filas duplicadas adicionales: {resumen['duplicados']}")
    for formato, cantidad in resumen["formatos"].items():
        print(f"Fechas con formato {formato}: {cantidad}")
    print(f"Fechas inválidas: {resumen['fechas_invalidas']}")
    print(f"Nombres vacíos: {resumen['nombres_vacios']}")
    print(f"Nombres con espacios sobrantes: {resumen['nombres_con_espacios']}")
    print(f"Nombres completamente en minúsculas: {resumen['nombres_minusculas']}")
    print(f"Nombres completamente en mayúsculas: {resumen['nombres_mayusculas']}")
    print(f"Archivo generado: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()
