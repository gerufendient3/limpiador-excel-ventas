"""Limpia un Excel de ventas y genera un informe de incidencias."""

from __future__ import annotations

import argparse
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNAS_REQUERIDAS = ["nombre", "fecha", "monto", "ciudad"]
FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y")
# Palabras que en nombres propios del español se escriben en minúscula.
PARTICULAS = {"de", "del", "la", "las", "los", "y", "e"}


def es_vacio(valor: Any) -> bool:
    """Reconoce celdas vacías sin confundir el número cero con un vacío."""
    if valor is None:
        return True
    try:
        return bool(pd.isna(valor))
    except (TypeError, ValueError):
        return False


def valor_informe(valor: Any) -> str:
    """Representación legible del valor original para el informe."""
    if es_vacio(valor):
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return str(valor)


def convertir_fecha(valor: Any) -> str | None:
    """Convierte solo formatos admitidos y fechas reales; no adivina fechas."""
    if es_vacio(valor):
        return None
    if isinstance(valor, (datetime, date, pd.Timestamp)):
        try:
            return valor.date().isoformat() if isinstance(valor, (datetime, pd.Timestamp)) else valor.isoformat()
        except (ValueError, OverflowError):
            return None

    texto = str(valor).strip()
    for formato in FORMATOS_FECHA:
        # strptime valida también que el día exista en el calendario.
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return None


def normalizar_nombre(texto: str) -> str:
    """Capitaliza cada palabra y deja en minúscula las partículas (de, la, del...)."""
    palabras = texto.split()
    resultado = []
    for i, palabra in enumerate(palabras):
        minuscula = palabra.lower()
        if i > 0 and minuscula in PARTICULAS:
            resultado.append(minuscula)
        else:
            resultado.append(minuscula.capitalize())
    return " ".join(resultado)


def es_numero(valor: Any) -> bool:
    if isinstance(valor, bool) or es_vacio(valor):
        return False
    if isinstance(valor, (int, float)):
        return math.isfinite(float(valor))
    if isinstance(valor, str):
        texto = valor.strip().replace("S/", "").replace("s/", "").strip().replace(",", "")
        if not texto:
            return False
        try:
            return math.isfinite(float(texto))
        except ValueError:
            return False
    return False


def numero(valor: Any) -> int | float | None:
    if not es_numero(valor):
        return None
    if isinstance(valor, str):
        texto = valor.strip().replace("S/", "").replace("s/", "").strip().replace(",", "")
        resultado = float(texto)
    else:
        resultado = float(valor)
    return int(resultado) if resultado.is_integer() else resultado


def registrar(incidencias: list[dict[str, Any]], fila: int, columna: str,
              tipo: str, valor: Any) -> None:
    incidencias.append({
        "fila_original": fila,
        "columna": columna,
        "tipo_error": tipo,
        "valor_original": valor_informe(valor),
    })


def procesar_archivo(entrada: Path) -> tuple[Path, Path, dict[str, int]]:
    if not entrada.is_file():
        raise FileNotFoundError(f"No existe el archivo de entrada: {entrada}")

    # Leer con openpyxl/pandas sin escribir nunca sobre el original.
    try:
        with pd.ExcelFile(entrada, engine="openpyxl") as libro:
            if "ventas" not in libro.sheet_names:
                raise ValueError("El archivo no contiene la hoja obligatoria 'ventas'.")
            df = pd.read_excel(libro, sheet_name="ventas", dtype=object, keep_default_na=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"No se pudo abrir el archivo Excel: {exc}") from exc
    faltantes = [col for col in COLUMNAS_REQUERIDAS if col not in df.columns]
    if faltantes:
        raise ValueError("Faltan columnas obligatorias en 'ventas': " + ", ".join(faltantes))
    df = df[COLUMNAS_REQUERIDAS].copy()
    filas_leidas = len(df)
    incidencias: list[dict[str, Any]] = []

    # Detectar duplicados sobre los valores originales, antes de limpiar espacios.
    indices_duplicados = df.duplicated(subset=COLUMNAS_REQUERIDAS, keep="first")
    for indice in df.index[indices_duplicados]:
        fila_excel = int(indice) + 2
        valor_fila = " | ".join(f"{col}={valor_informe(df.at[indice, col])}" for col in COLUMNAS_REQUERIDAS)
        registrar(incidencias, fila_excel, "fila completa", "duplicado exacto eliminado", valor_fila)
    df = df.loc[~indices_duplicados].copy()

    fechas_convertidas = 0
    fechas_invalidas = 0
    nombres_vacios = 0
    montos_error = 0

    for indice in df.index:
        fila_excel = int(indice) + 2

        # Recortar todos los textos; la comparación de duplicados ya se hizo.
        for columna in ("nombre", "fecha", "ciudad"):
            valor = df.at[indice, columna]
            if isinstance(valor, str):
                df.at[indice, columna] = valor.strip()

        nombre = df.at[indice, "nombre"]
        if es_vacio(nombre) or (isinstance(nombre, str) and not nombre.strip()):
            nombres_vacios += 1
            registrar(incidencias, fila_excel, "nombre", "nombre vacío", nombre)
            df.at[indice, "nombre"] = ""
        else:
            df.at[indice, "nombre"] = normalizar_nombre(str(nombre))

        fecha_original = df.at[indice, "fecha"]
        fecha_limpia = convertir_fecha(fecha_original)
        if fecha_limpia is None:
            # Una celda de fecha vacía ya estaba vacía; solo reportamos valores no vacíos inválidos.
            if not es_vacio(fecha_original) and str(fecha_original).strip():
                fechas_invalidas += 1
                registrar(incidencias, fila_excel, "fecha", "fecha inválida", fecha_original)
            df.at[indice, "fecha"] = ""
        else:
            if str(fecha_original).strip() != fecha_limpia:
                fechas_convertidas += 1
            df.at[indice, "fecha"] = fecha_limpia

        monto_original = df.at[indice, "monto"]
        if not es_numero(monto_original):
            montos_error += 1
            registrar(incidencias, fila_excel, "monto", "monto no numérico", monto_original)
            df.at[indice, "monto"] = ""
        else:
            monto_limpio = numero(monto_original)
            df.at[indice, "monto"] = monto_limpio
            if monto_limpio < 0:
                montos_error += 1
                registrar(incidencias, fila_excel, "monto", "monto negativo", monto_original)

    salida_limpia = entrada.with_name("ventas_limpias.xlsx")
    salida_informe = entrada.with_name("informe_errores.xlsx")
    guardar_ventas(df, salida_limpia)
    guardar_informe(incidencias, salida_informe)

    resumen = {
        "filas_leidas": filas_leidas,
        "duplicados_eliminados": int(indices_duplicados.sum()),
        "filas_finales": len(df),
        "fechas_convertidas": fechas_convertidas,
        "fechas_invalidas": fechas_invalidas,
        "nombres_vacios": nombres_vacios,
        "montos_con_error": montos_error,
        "incidencias_totales": len(incidencias),
    }
    return salida_limpia, salida_informe, resumen


def dar_formato(hoja) -> None:
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="24527A")
        celda.alignment = Alignment(horizontal="center")
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        hoja.column_dimensions[letra].width = min(max(ancho + 2, 12), 48)


def guardar_ventas(df: pd.DataFrame, ruta: Path) -> None:
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="ventas")
        dar_formato(writer.sheets["ventas"])


def guardar_informe(incidencias: list[dict[str, Any]], ruta: Path) -> None:
    columnas = ["fila_original", "columna", "tipo_error", "valor_original"]
    df_inc = pd.DataFrame(incidencias, columns=columnas)
    if df_inc.empty:
        df_resumen = pd.DataFrame(columns=["tipo_error", "cantidad"])
    else:
        df_resumen = (df_inc.groupby("tipo_error", dropna=False).size()
                      .rename("cantidad").reset_index().sort_values("tipo_error"))
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df_inc.to_excel(writer, index=False, sheet_name="incidencias")
        df_resumen.to_excel(writer, index=False, sheet_name="resumen")
        for hoja in writer.sheets.values():
            dar_formato(hoja)


def main() -> int:
    parser = argparse.ArgumentParser(description="Limpia un Excel de ventas sin modificar el original.")
    parser.add_argument("archivo", type=Path, help="Ruta del Excel de entrada")
    args = parser.parse_args()
    try:
        limpia, informe, resumen = procesar_archivo(args.archivo.resolve())
    except (FileNotFoundError, ValueError, PermissionError) as exc:
        parser.exit(2, f"Error: {exc}\n")
    except Exception as exc:
        parser.exit(1, f"Error inesperado al procesar el Excel: {exc}\n")

    print("LIMPIEZA COMPLETADA")
    print(f"Filas leídas: {resumen['filas_leidas']}")
    print(f"Duplicados exactos eliminados: {resumen['duplicados_eliminados']}")
    print(f"Filas en el Excel limpio: {resumen['filas_finales']}")
    print(f"Fechas convertidas: {resumen['fechas_convertidas']}")
    print(f"Fechas inválidas: {resumen['fechas_invalidas']}")
    print(f"Nombres vacíos: {resumen['nombres_vacios']}")
    print(f"Montos con error: {resumen['montos_con_error']}")
    print(f"Incidencias totales: {resumen['incidencias_totales']}")
    print(f"Excel limpio: {limpia}")
    print(f"Informe de errores: {informe}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
